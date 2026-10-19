"""Extract workbook cells, formulas, and regions with OOXML safety checks."""

from __future__ import annotations

import re
import xml.etree.ElementTree as ET
from dataclasses import replace
from importlib.metadata import version as package_version
from pathlib import Path, PurePosixPath
from typing import cast
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

from ..storage import source_sha256
from .spreadsheet_models import (
    SemanticWorkbook,
    SourceType,
    WorkbookCell,
    WorkbookChart,
    WorkbookRegion,
    WorkbookSheet,
)

MAX_OOXML_ENTRIES = 20_000
MAX_OOXML_UNCOMPRESSED_BYTES = 512 * 1024 * 1024
MAX_COMPRESSION_RATIO = 200
MAX_WORKBOOK_SHEETS = 250
MAX_WORKBOOK_CELLS = 2_000_000
_FORMULA_REFERENCE = re.compile(
    r"(?:'([^']+)'|([A-Za-z_][\w .-]*))!\$?[A-Z]{1,3}\$?\d+"
)


class SpreadsheetConversionError(RuntimeError):
    pass


class SpreadsheetConverter:
    def convert(self, path: Path) -> SemanticWorkbook:
        """Return workbook evidence after validating the XLSX archive.

        Formulas are retained with available cached values and never executed.
        """
        security_warnings = _validate_ooxml(path)
        formula_book = None
        try:
            formula_book = load_workbook(
                path, data_only=False, read_only=False, keep_links=False
            )
            value_book = load_workbook(
                path, data_only=True, read_only=False, keep_links=False
            )
        except Exception as error:
            if formula_book is not None:
                formula_book.close()
            raise SpreadsheetConversionError(
                "The XLSX workbook could not be read safely."
            ) from error
        warnings: list[str] = list(security_warnings)
        sheets: list[WorkbookSheet] = []
        try:
            if len(formula_book.worksheets) > MAX_WORKBOOK_SHEETS:
                raise SpreadsheetConversionError(
                    "XLSX contains too many worksheets to process safely."
                )
            estimated_cells = sum(
                sheet.max_row * sheet.max_column for sheet in formula_book.worksheets
            )
            if estimated_cells > MAX_WORKBOOK_CELLS:
                raise SpreadsheetConversionError(
                    "XLSX contains too many cells to process safely."
                )
            calculation = getattr(formula_book, "calculation", None)
            if calculation is not None and (
                getattr(calculation, "calcMode", None) == "manual"
                or bool(getattr(calculation, "fullCalcOnLoad", False))
            ):
                warnings.append("formula_cache_potentially_stale")
            named_regions = _named_regions(formula_book)
            for sheet in formula_book.worksheets:
                value_sheet = value_book[sheet.title]
                sheets.append(
                    _sheet_model(
                        sheet,
                        value_sheet,
                        warnings,
                        named_regions.get(sheet.title, ()),
                    )
                )
            named_ranges = _named_ranges(formula_book)
        finally:
            formula_book.close()
            value_book.close()
        return SemanticWorkbook(
            path.name,
            SourceType.XLSX,
            tuple(sheets),
            named_ranges,
            tuple(dict.fromkeys(warnings)),
            package_version("openpyxl"),
            {"source_size": path.stat().st_size, "source_sha256": source_sha256(path)},
        )


def _validate_ooxml(path: Path) -> tuple[str, ...]:
    if path.suffix.casefold() != ".xlsx":
        raise SpreadsheetConversionError("Only .xlsx workbooks are supported.")
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_OOXML_ENTRIES:
                raise SpreadsheetConversionError("XLSX contains too many ZIP entries.")
            total = 0
            names = set()
            for entry in entries:
                member = PurePosixPath(entry.filename)
                if (
                    member.is_absolute()
                    or ".." in member.parts
                    or "\\" in entry.filename
                ):
                    raise SpreadsheetConversionError(
                        "XLSX contains an unsafe ZIP path."
                    )
                total += entry.file_size
                if total > MAX_OOXML_UNCOMPRESSED_BYTES:
                    raise SpreadsheetConversionError(
                        "XLSX expands beyond the safe size limit."
                    )
                if (
                    entry.file_size > 1_000_000
                    and entry.file_size / max(1, entry.compress_size)
                    > MAX_COMPRESSION_RATIO
                ):
                    raise SpreadsheetConversionError(
                        "XLSX has a suspicious compression ratio."
                    )
                names.add(entry.filename)
            if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
                raise SpreadsheetConversionError("File is not a valid XLSX workbook.")
            if any(name.casefold().endswith("vbaproject.bin") for name in names):
                raise SpreadsheetConversionError(
                    "Macro-enabled workbooks are not supported."
                )
            warnings: list[str] = []
            for relationship in (
                name for name in names if name.casefold().endswith(".rels")
            ):
                if archive.getinfo(relationship).file_size > 2 * 1024 * 1024:
                    raise SpreadsheetConversionError(
                        "XLSX relationship metadata is too large."
                    )
                content = archive.read(relationship)
                try:
                    root = ET.fromstring(content)
                except ET.ParseError as error:
                    raise SpreadsheetConversionError(
                        "XLSX contains malformed relationship metadata."
                    ) from error
                for item in root:
                    if item.attrib.get("TargetMode", "").casefold() != "external":
                        continue
                    relationship_type = item.attrib.get("Type", "").casefold()
                    if any(
                        unsafe in relationship_type
                        for unsafe in (
                            "externallink",
                            "oleobject",
                            "attachedtemplate",
                            "package",
                        )
                    ):
                        raise SpreadsheetConversionError(
                            "XLSX contains an unsafe external data or embedded-object relationship."
                        )
                    warnings.append("external_hyperlink_retained_without_fetch")
            worksheet_names = [
                name
                for name in names
                if name.startswith("xl/worksheets/") and name.endswith(".xml")
            ]
            if len(worksheet_names) > MAX_WORKBOOK_SHEETS:
                raise SpreadsheetConversionError(
                    "XLSX contains too many worksheets to process safely."
                )
            cell_count = 0
            for worksheet_name in worksheet_names:
                with archive.open(worksheet_name) as worksheet:
                    try:
                        for _, element in ET.iterparse(worksheet, events=("end",)):
                            if element.tag.rsplit("}", 1)[-1] == "c":
                                cell_count += 1
                                if cell_count > MAX_WORKBOOK_CELLS:
                                    raise SpreadsheetConversionError(
                                        "XLSX contains too many cells to process safely."
                                    )
                            element.clear()
                    except ET.ParseError as error:
                        raise SpreadsheetConversionError(
                            "XLSX contains malformed worksheet XML."
                        ) from error
            return tuple(dict.fromkeys(warnings))
    except BadZipFile as error:
        raise SpreadsheetConversionError(
            "File is not a valid XLSX ZIP container."
        ) from error


def _sheet_model(
    sheet, value_sheet, warnings: list[str], named_regions=()
) -> WorkbookSheet:
    merged = tuple(str(item) for item in sheet.merged_cells.ranges)
    merged_by_cell: dict[str, str] = {}
    for value in merged:
        min_col, min_row, max_col, max_row = cast(
            tuple[int, int, int, int], range_boundaries(value)
        )
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                merged_by_cell[f"{get_column_letter(column)}{row}"] = value
    hidden_rows = tuple(
        index for index, dimension in sheet.row_dimensions.items() if dimension.hidden
    )
    hidden_columns = tuple(
        key for key, dimension in sheet.column_dimensions.items() if dimension.hidden
    )
    cells: list[WorkbookCell] = []
    # Serialized cells distinguish explicit blanks from coordinates absent in the source.
    for cell in sorted(
        sheet._cells.values(), key=lambda value: (value.row, value.column)
    ):
        cached = value_sheet[cell.coordinate].value
        formula = str(cell.value) if cell.data_type == "f" else None
        if formula and cached is None:
            warnings.append(f"missing_formula_cache:{sheet.title}!{cell.coordinate}")
        cells.append(
            WorkbookCell(
                coordinate=cell.coordinate,
                row=cell.row,
                column=cell.column,
                value=cached if formula else cell.value,
                formula=formula,
                cached_value=cached if formula else None,
                value_type=_value_type(cell, cached),
                number_format=cell.number_format
                if cell.number_format != "General"
                else None,
                currency_or_unit=_currency_or_unit(cell.number_format),
                bold=bool(cell.font and cell.font.bold),
                italic=bool(cell.font and cell.font.italic),
                fill_color=_fill_color(cell),
                indent=int(cell.alignment.indent or 0),
                hidden=cell.row in hidden_rows
                or get_column_letter(cell.column) in hidden_columns
                or sheet.sheet_state != "visible",
                comment=cell.comment.text if cell.comment else None,
                merged_range=merged_by_cell.get(cell.coordinate),
                cross_sheet_references=tuple(
                    dict.fromkeys(_formula_references(formula))
                ),
                raw_value=cell.value,
                blank_kind="empty_cell" if cell.value is None else None,
            )
        )
    regions = _regions(sheet, cells, named_regions)
    charts = tuple(
        _chart_model(chart, index)
        for index, chart in enumerate(_worksheet_charts(sheet, warnings), 1)
    )
    if getattr(sheet, "_images", ()):
        warnings.append(f"sheet_images_not_extracted:{sheet.title}")
    return WorkbookSheet(
        sheet.title,
        sheet.sheet_state,
        tuple(cells),
        regions,
        merged,
        charts,
        hidden_rows,
        hidden_columns,
    )


def _regions(
    sheet, cells: list[WorkbookCell], named_regions=()
) -> tuple[WorkbookRegion, ...]:
    regions: list[WorkbookRegion] = []
    occupied: set[tuple[int, int]] = set()
    for table in sheet.tables.values():
        min_col, min_row, max_col, max_row = cast(
            tuple[int, int, int, int], range_boundaries(table.ref)
        )
        table_cells = {
            (row, column)
            for row in range(min_row, max_row + 1)
            for column in range(min_col, max_col + 1)
        }
        regions.append(
            WorkbookRegion(
                f"region-{len(regions) + 1:03d}",
                table.displayName,
                table.ref,
                "excel_table",
                min_row,
                max_row,
                min_col,
                max_col,
                (min_row,),
            )
        )
        occupied.update(table_cells)
    for original in named_regions:
        region_cells = {
            (row, column)
            for row in range(original.min_row, original.max_row + 1)
            for column in range(original.min_column, original.max_column + 1)
        }
        warning = "overlaps_preferred_region" if region_cells & occupied else None
        regions.append(
            replace(
                original,
                id=f"region-{len(regions) + 1:03d}",
                warnings=tuple(
                    dict.fromkeys(
                        (*original.warnings, *((warning,) if warning else ()))
                    )
                ),
            )
        )
        occupied.update(region_cells)
    coordinates = {
        (cell.row, cell.column)
        for cell in cells
        if (cell.row, cell.column) not in occupied
    }
    while coordinates:
        seed = min(coordinates)
        component = {seed}
        frontier = [seed]
        while frontier:
            row, column = frontier.pop()
            for neighbor in (
                (row - 1, column),
                (row + 1, column),
                (row, column - 1),
                (row, column + 1),
            ):
                if neighbor in coordinates and neighbor not in component:
                    component.add(neighbor)
                    frontier.append(neighbor)
        coordinates -= component
        min_row = min(row for row, _ in component)
        max_row = max(row for row, _ in component)
        min_col = min(column for _, column in component)
        max_col = max(column for _, column in component)
        ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
        warnings = (
            ("sparse_or_ambiguous_region",)
            if len(component) < (max_row - min_row + 1) * (max_col - min_col + 1)
            else ()
        )
        regions.append(
            WorkbookRegion(
                f"region-{len(regions) + 1:03d}",
                f"Region {len(regions) + 1}",
                ref,
                min_row=min_row,
                max_row=max_row,
                min_column=min_col,
                max_column=max_col,
                header_rows=(),
                warnings=warnings,
                kind="inferred",
            )
        )
    return tuple(_classify_region(region, cells) for region in regions)


def _classify_region(
    region: WorkbookRegion, cells: list[WorkbookCell]
) -> WorkbookRegion:
    selected = [
        cell
        for cell in cells
        if region.min_row <= cell.row <= region.max_row
        and region.min_column <= cell.column <= region.max_column
    ]
    by_row: dict[int, list[WorkbookCell]] = {}
    for cell in selected:
        by_row.setdefault(cell.row, []).append(cell)
    title_rows: list[int] = []
    header_rows: list[int] = []
    total_rows: list[int] = []
    subtotal_rows: list[int] = []
    note_rows: list[int] = []
    for row in range(region.min_row, region.max_row + 1):
        row_cells = by_row.get(row, [])
        text = " ".join(str(cell.value or "") for cell in row_cells).casefold()
        if (
            row == region.min_row
            and len(row_cells) <= 2
            and any(cell.bold or cell.merged_range for cell in row_cells)
        ):
            title_rows.append(row)
        elif row_cells and sum(cell.bold for cell in row_cells) >= max(
            1, len(row_cells) // 2
        ):
            header_rows.append(row)
        if re.search(r"\bgrand total\b|\btotal\b", text):
            total_rows.append(row)
        elif re.search(r"\bsubtotal\b", text):
            subtotal_rows.append(row)
        if re.search(r"\bnote\b|\bsource\b|\bassumption\b", text):
            note_rows.append(row)
    width = region.max_column - region.min_column + 1
    parameter_block = (
        width <= 3
        and len(selected) >= 2
        and sum(cell.bold for cell in selected if cell.column == region.min_column) >= 1
    )
    role = (
        "parameters"
        if parameter_block
        else "notes"
        if note_rows and len(note_rows) >= max(1, len(by_row) // 2)
        else "table"
    )
    source = "native_table" if region.kind == "excel_table" else "inferred"
    confidence = 0.98 if source == "native_table" else 0.72
    ambiguous = bool(region.warnings)
    if ambiguous:
        confidence = min(confidence, 0.55)
    return replace(
        region,
        header_rows=tuple(header_rows or region.header_rows),
        semantic_role=role,
        confidence=confidence,
        source=source,
        title_rows=tuple(title_rows),
        total_rows=tuple(total_rows),
        subtotal_rows=tuple(subtotal_rows),
        note_rows=tuple(note_rows),
        parameter_block=parameter_block,
        ambiguous=ambiguous,
    )


def _named_ranges(book) -> tuple[dict[str, object], ...]:
    result = []
    for name, definition in book.defined_names.items():
        destinations = []
        try:
            destinations = [
                {"sheet": sheet, "range": reference}
                for sheet, reference in definition.destinations
            ]
        except (AttributeError, TypeError):
            pass
        result.append(
            {
                "name": name,
                "destinations": destinations,
                "hidden": bool(getattr(definition, "hidden", False)),
            }
        )
    return tuple(result)


def _named_regions(book) -> dict[str, tuple[WorkbookRegion, ...]]:
    result: dict[str, list[WorkbookRegion]] = {}
    for name, definition in book.defined_names.items():
        try:
            destinations = list(definition.destinations)
        except (AttributeError, TypeError):
            continue
        for sheet_name, reference in destinations:
            try:
                min_col, min_row, max_col, max_row = cast(
                    tuple[int, int, int, int],
                    range_boundaries(reference.replace("$", "")),
                )
            except (TypeError, ValueError):
                continue
            items = result.setdefault(sheet_name, [])
            items.append(
                WorkbookRegion(
                    f"named-{len(items) + 1:03d}",
                    name,
                    reference.replace("$", ""),
                    "named_range",
                    min_row,
                    max_row,
                    min_col,
                    max_col,
                )
            )
    return {sheet: tuple(items) for sheet, items in result.items()}


def _chart_model(chart, index: int) -> WorkbookChart:
    series = []
    for item in getattr(chart, "ser", ()):
        series.append(
            {
                "title": str(getattr(item, "tx", "") or "") or None,
                "categories": _series_reference(getattr(item, "cat", None)),
                "values": _series_reference(getattr(item, "val", None)),
            }
        )
    return WorkbookChart(
        f"chart-{index:03d}",
        _chart_title(chart),
        type(chart).__name__,
        _anchor_coordinate(chart),
        tuple(series),
    )


def _worksheet_charts(sheet, warnings: list[str]) -> tuple[object, ...]:
    charts = getattr(sheet, "_charts", None)
    if charts is None:
        warnings.append(f"chart_metadata_unavailable:{sheet.title}")
        return ()
    return tuple(charts)


def _anchor_coordinate(value) -> str | None:
    anchor = getattr(value, "anchor", None)
    if isinstance(anchor, str):
        return anchor
    marker = getattr(anchor, "_from", None)
    if marker is not None:
        return f"{get_column_letter(marker.col + 1)}{marker.row + 1}"
    return None


def _chart_title(chart) -> str | None:
    title = getattr(chart, "title", None)
    if title is None:
        return None
    try:
        paragraphs = title.tx.rich.p
        text = " ".join(
            run.t for paragraph in paragraphs for run in paragraph.r if run.t
        )
        return text or None
    except AttributeError:
        return None


def _series_reference(value) -> str | None:
    for name in ("numRef", "strRef"):
        reference = getattr(value, name, None)
        if reference is not None and getattr(reference, "f", None):
            return reference.f
    return None


def _formula_references(formula: str | None) -> list[str]:
    if not formula:
        return []
    return [
        (match.group(1) or match.group(2)).strip()
        for match in _FORMULA_REFERENCE.finditer(formula)
    ]


def _value_type(cell, value) -> str:
    if getattr(cell, "is_date", False):
        return "date"
    if cell.data_type == "f":
        return "formula"
    if isinstance(value if value is not None else cell.value, bool):
        return "boolean"
    if isinstance(value if value is not None else cell.value, (int, float)):
        return "number"
    if cell.data_type == "e":
        return "error"
    return "text"


def _currency_or_unit(number_format: str) -> str | None:
    if "%" in number_format:
        return "percent"
    for symbol, name in (("$", "USD"), ("€", "EUR"), ("£", "GBP"), ("¥", "JPY")):
        if symbol in number_format:
            return name
    return None


def _fill_color(cell) -> str | None:
    color = getattr(cell.fill, "fgColor", None)
    rgb = getattr(color, "rgb", None)
    return rgb if isinstance(rgb, str) and rgb not in {"00000000", "000000"} else None
