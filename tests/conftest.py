from datetime import datetime

import pymupdf
import pytest
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "boq.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "BOQ"
    sheet.append(["Concrete works"])
    sheet.merge_cells("A1:F1")
    sheet["A1"].font = Font(bold=True)
    sheet.append(["Item", "Description", "Unit", "Qty", "Rate", "Amount"])
    sheet.append(["001", "Concrete – أساس", "m3", 0, 12.5, "=D3*E3"])
    sheet.append(["002", "Rebar", "kg", -2.5, 10, -25])
    sheet["E3"].number_format = "$0.00"
    sheet.row_dimensions[4].hidden = True
    sheet.add_table(Table(displayName="BOQItems", ref="A2:F4"))
    other = book.create_sheet("Notes")
    other.append([datetime(2026, 10, 19), False, "=BOQ!F3"])
    book.save(path)
    book.close()
    return path


@pytest.fixture
def pdf(tmp_path):
    path = tmp_path / "spec.pdf"
    with pymupdf.open() as document:
        page = document.new_page()
        page.insert_text((40, 40), "1 Concrete", fontsize=20)
        page.insert_text(
            (40, 70), "Strength is 35 MPa; tolerance -2.5 mm.", fontsize=10
        )
        xs = [40, 90, 220, 270, 320, 380, 450]
        ys = [100, 130, 160, 190]
        for x in xs:
            page.draw_line((x, ys[0]), (x, ys[-1]))
        for y in ys:
            page.draw_line((xs[0], y), (xs[-1], y))
        for row, values in enumerate(
            [
                ["Item", "Description", "Unit", "Qty", "Rate", "Amount"],
                ["001", "Concrete", "m3", "0", "12.50", "0.00"],
                ["002", "Rebar", "kg", "-2.5", "10", "-25"],
            ]
        ):
            for col, value in enumerate(values):
                page.insert_text((xs[col] + 3, ys[row] + 19), value, fontsize=9)
        page = document.new_page()
        page.insert_text((40, 40), "2 Installation", fontsize=20)
        page.insert_text((40, 70), "Install at 100 mm centres.", fontsize=10)
        document.set_toc([[1, "1 Concrete", 1], [1, "2 Installation", 2]])
        document.save(path)
    return path
